from copy import deepcopy
from pathlib import Path
import hashlib
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\CV_Android_NguyenPhiVu.docx")
OUTPUT = Path(r"D:\frontend\artifacts\CV_Java_Fresher_NguyenPhiVu_PulseTech.docx")
EXPECTED_SHA256 = "B7E093672BC27F12ABD94CD8C093CEED03FC5465F404ED419590F03DCD3F0F6B"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def clear_content(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def add_text(paragraph, text, bold=None, italic=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    return run


def set_plain(paragraph, text, *, bold=None, italic=None, size=None):
    clear_content(paragraph)
    add_text(paragraph, text, bold=bold, italic=italic, size=size)


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_labeled_links(paragraph, label, links):
    clear_content(paragraph)
    add_text(paragraph, label, bold=True)
    add_text(paragraph, " ")
    for index, (text, url) in enumerate(links):
        if index:
            add_text(paragraph, "  |  ")
        add_hyperlink(paragraph, text, url)


def clone_bullet_before(anchor, source_bullet, text):
    new_paragraph = anchor.insert_paragraph_before(text, style=source_bullet.style)
    copy_numbering(new_paragraph, source_bullet)
    p_pr = new_paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "35")
    spacing.set(qn("w:after"), "0")
    return new_paragraph


def copy_numbering(target, source):
    target_pr = target._p.get_or_add_pPr()
    source_pr = source._p.get_or_add_pPr()
    existing = target_pr.find(qn("w:numPr"))
    if existing is not None:
        target_pr.remove(existing)
    source_num = source_pr.find(qn("w:numPr"))
    if source_num is not None:
        target_pr.append(deepcopy(source_num))


actual_hash = sha256(SOURCE)
if actual_hash != EXPECTED_SHA256:
    raise RuntimeError(f"Template hash changed: {actual_hash}")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)
paragraphs = doc.paragraphs

# Identity and positioning
set_plain(paragraphs[0], "NGUYEN PHI VU", bold=True, size=16)
set_plain(paragraphs[1], "Fresher Java Backend Developer", bold=True, size=13)

# Career objective
set_plain(paragraphs[2], "CAREER OBJECTIVE", bold=True, size=14)
set_plain(
    paragraphs[3],
    "Final-year Software Engineering student with hands-on experience building a full-stack e-commerce platform "
    "with Java 21, Spring Boot, Spring Cloud Gateway, MongoDB, and Docker. Seeking a Fresher Java Backend Developer "
    "position to strengthen backend engineering skills, contribute to reliable business systems, and grow through "
    "code review and collaboration with an experienced development team.",
)

# Education stays factual and keeps the source layout
set_plain(paragraphs[4], "EDUCATION", bold=True, size=14)
clear_content(paragraphs[5])
add_text(paragraphs[5], "Industrial University of Ho Chi Minh City", bold=True)
add_text(paragraphs[5], "\t\t\t\t\t\t     September 2021 - Present")
add_text(paragraphs[5], "\nBachelor of Software Engineering", italic=True)
add_text(paragraphs[5], "\n\nSKILLS", bold=True, size=14)

# Skills: preserve the source's alternating one-column and two-column sections
set_plain(paragraphs[6], "Java & Backend Development", bold=True)
skills_backend = {
    8: "Programming Languages: Java 21, JavaScript, TypeScript.",
    9: "Frameworks: Spring Boot 4.1, Spring Cloud Gateway, Spring Data MongoDB.",
    10: "Architecture: Microservices, RESTful API, API Gateway, inter-service communication.",
    11: "Database: MongoDB Atlas; document-oriented data modeling.",
    12: "Authentication: Email verification, JWT-based admin access, role-based authorization.",
    13: "Payment Integration: VNPay, MoMo, Stripe, and Cash on Delivery.",
    14: "Frontend Knowledge: Next.js 16, React 19, TypeScript, Tailwind CSS 4.",
}
for index, text in skills_backend.items():
    set_plain(paragraphs[index], text)

set_plain(paragraphs[15], "DevOps & Infrastructure", bold=True)
skills_devops = {
    16: "Containerization: Docker, Docker Compose.",
    17: "Deployment: Vercel and Render.",
    18: "Build Tools: Maven multi-module, npm.",
    19: "Version Control: Git, GitHub.",
    20: "API Integration: REST clients, callback and webhook handling.",
    21: "Testing: JUnit and Spring Boot Test.",
}
for index, text in skills_devops.items():
    set_plain(paragraphs[index], text)
copy_numbering(paragraphs[20], paragraphs[21])

set_plain(paragraphs[22], "Development Tools", bold=True)
set_plain(paragraphs[24], "IDE: IntelliJ IDEA, Visual Studio Code.")
set_plain(paragraphs[25], "API Testing: Postman and browser network tools.")
set_plain(paragraphs[26], "Other: Maven, validation, SMTP email, Cloudinary, Git workflows.")

# Project section
set_plain(paragraphs[27], "\nPERSONAL PROJECTS", bold=True, size=13)
set_plain(paragraphs[28], "PulseTech E-Commerce Platform", bold=True, size=12)
set_labeled_links(
    paragraphs[29],
    "Source:",
    [
        ("Frontend", "https://github.com/VuIceTea/PulseTech"),
        ("Admin", "https://github.com/VuIceTea/PulseTech-Admin"),
        ("Backend", "https://github.com/VuIceTea/PulseTech-Backend"),
    ],
)
set_labeled_links(
    paragraphs[30],
    "Live Demo:",
    [("pulse-tech-beryl.vercel.app", "https://pulse-tech-beryl.vercel.app")],
)

project_bullets = [
    "Designed and developed a technology e-commerce system with a Next.js storefront, a separate admin dashboard, and Java Spring Boot microservices.",
    "Implemented four Maven modules: API Gateway, Product Service, Auth Service, and Order Service, with Spring Cloud Gateway as the single entry point.",
    "Built REST APIs for products, variants, content, authentication, addresses, wishlists, carts, coupons, orders, and payment workflows.",
    "Modeled color and storage variants with independent price offsets and stock; validated availability and decreased the correct variant stock during ordering.",
    "Integrated VNPay, MoMo, Stripe, and COD, including payment URL creation, return callbacks, signature verification, and order payment updates.",
    "Developed coupon rules for product and shipping discounts, customer assignment, usage limits, and one coupon per discount category.",
    "Built admin features for product and inventory management, order details and status transitions, vouchers, customer locking, and role assignment.",
    "Integrated Gemini AI with live product, variant-stock, pricing, and policy context to provide data-grounded customer support.",
    "Containerized the system with Docker multi-stage builds and Docker Compose; configured MongoDB Atlas and environment-based deployment for Vercel and Render.",
]
for index, text in zip(range(31, 36), project_bullets[:5]):
    set_plain(paragraphs[index], text)

anchor = paragraphs[36]
source_bullet = paragraphs[35]
for text in project_bullets[5:]:
    clone_bullet_before(anchor, source_bullet, text)

# Contact section
set_plain(paragraphs[36], "CONTACT", bold=True, size=13)
clear_content(paragraphs[37])
add_text(paragraphs[37], "Email:", bold=True, size=13)
add_text(paragraphs[37], " ", size=13)
add_hyperlink(paragraphs[37], "zozotrong123456@gmail.com", "mailto:zozotrong123456@gmail.com")
add_text(paragraphs[37], "\nPhone:", bold=True, size=13)
add_text(paragraphs[37], " 0329982474", size=13)
add_text(paragraphs[37], "\nGitHub:", bold=True, size=13)
add_text(paragraphs[37], " ", size=13)
add_hyperlink(paragraphs[37], "https://github.com/VuIceTea", "https://github.com/VuIceTea")
add_text(paragraphs[37], "\nLinkedIn:", bold=True, size=13)
add_text(paragraphs[37], " ", size=13)
add_hyperlink(
    paragraphs[37],
    "https://www.linkedin.com/in/phi-vu-nguyen-8a18a5371/",
    "https://www.linkedin.com/in/phi-vu-nguyen-8a18a5371/",
)

doc.core_properties.title = "Nguyen Phi Vu Fresher Java Backend Developer CV"
doc.core_properties.subject = "PulseTech Java Backend project CV"
doc.core_properties.author = "Nguyen Phi Vu"
doc.core_properties.keywords = "Java, Spring Boot, Microservices, MongoDB, Docker, Fresher Java Backend"
doc.save(OUTPUT)

print(OUTPUT)
print("source_sha256=" + sha256(SOURCE))
print("output_sha256=" + sha256(OUTPUT))
