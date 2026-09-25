from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import zipfile

path = Path(r"D:\CV_Android_NguyenPhiVu.docx")
doc = Document(path)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for i, p in enumerate(doc.paragraphs):
    pf = p.paragraph_format
    runs = []
    for r in p.runs:
        runs.append({
            "text": r.text,
            "bold": r.bold,
            "italic": r.italic,
            "size": None if r.font.size is None else round(r.font.size.pt, 2),
            "name": r.font.name,
            "color": None if r.font.color is None or r.font.color.rgb is None else str(r.font.color.rgb),
        })
    print({
        "i": i,
        "text": p.text,
        "style": p.style.name,
        "alignment": p.alignment,
        "left_cm": None if pf.left_indent is None else round(pf.left_indent.cm, 2),
        "first_cm": None if pf.first_line_indent is None else round(pf.first_line_indent.cm, 2),
        "before_pt": None if pf.space_before is None else round(pf.space_before.pt, 2),
        "after_pt": None if pf.space_after is None else round(pf.space_after.pt, 2),
        "line": pf.line_spacing,
        "runs": runs,
    })

for i, s in enumerate(doc.sections, 1):
    cols = s._sectPr.find(qn("w:cols"))
    print("SECTION", i, "type", s.start_type, "cols", None if cols is None else dict(cols.attrib))

with zipfile.ZipFile(path) as z:
    print("PARTS")
    for info in z.infolist():
        print(info.filename, info.file_size)
