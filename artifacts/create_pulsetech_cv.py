from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).with_name("CV_Nguyen_Phi_Vu_PulseTech.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color="1F4E79"):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.append(fonts)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    keep_with_next(p)
    r = p.add_run(text)
    set_font(r, size=12.5, bold=True, color="000000")
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.04
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, size=9.7, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, size=9.7)
    else:
        r = p.add_run(text)
        set_font(r, size=9.7)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.15)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.7)
normal.font.color.rgb = RGBColor(31, 31, 31)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

for style_name in ("Heading 1", "Heading 2"):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
title.add_run("NGUYỄN PHI VŨ")

role = doc.add_paragraph()
role.alignment = WD_ALIGN_PARAGRAPH.CENTER
role.paragraph_format.space_after = Pt(3)
r = role.add_run("FULL STACK DEVELOPER")
set_font(r, size=12, bold=True, color="1F4E79")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(7)
r = contact.add_run("Việt Nam  |  ")
set_font(r, size=9.4, color="555555")
add_hyperlink(contact, "github.com/VuIceTea", "https://github.com/VuIceTea")
r = contact.add_run("  |  ")
set_font(r, size=9.4, color="555555")
add_hyperlink(contact, "PulseTech Demo", "https://pulse-tech-beryl.vercel.app")

add_section_heading(doc, "TÓM TẮT CHUYÊN MÔN")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.line_spacing = 1.08
r = p.add_run(
    "Full Stack Developer định hướng xây dựng ứng dụng web có nghiệp vụ hoàn chỉnh. "
    "Đã phát triển PulseTech, nền tảng thương mại điện tử gồm website khách hàng, trang quản trị "
    "và backend microservices. Có kinh nghiệm với Next.js, React, TypeScript, Java, Spring Boot, "
    "MongoDB, Docker, tích hợp thanh toán và trợ lý AI dựa trên dữ liệu hệ thống."
)
set_font(r, size=9.8)

add_section_heading(doc, "NĂNG LỰC KỸ THUẬT")
skills = [
    ("Frontend", "Next.js 16, React 19, TypeScript, Tailwind CSS 4, App Router, responsive UI"),
    ("Backend", "Java 21, Spring Boot 4.1, Spring Cloud Gateway, REST API, validation"),
    ("Dữ liệu", "MongoDB Atlas, Spring Data MongoDB, quản lý dữ liệu theo microservice"),
    ("Tích hợp", "VNPay, MoMo, Stripe, Gemini AI, Cloudinary, SMTP email"),
    ("DevOps", "Docker, Docker Compose, Maven multi-module, Vercel, Render"),
    ("Nghiệp vụ", "E-commerce, sản phẩm biến thể, tồn kho, voucher, thanh toán, vòng đời đơn hàng"),
]
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Cm(3.25)
table.columns[1].width = Cm(14.0)
hdr = table.rows[0].cells
hdr[0].text = "NHÓM"
hdr[1].text = "CÔNG NGHỆ VÀ KỸ NĂNG"
set_repeat_table_header(table.rows[0])
for cell in hdr:
    set_cell_shading(cell, "1F4E79")
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        set_font(run, size=9, bold=True, color="FFFFFF")
for index, (group, values) in enumerate(skills):
    cells = table.add_row().cells
    cells[0].text = group
    cells[1].text = values
    for cell in cells:
        set_cell_border(cell)
        set_cell_margins(cell, top=80, bottom=80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index % 2 == 1:
            set_cell_shading(cell, "F4F7FA")
    for run in cells[0].paragraphs[0].runs:
        set_font(run, size=9, bold=True)
    for run in cells[1].paragraphs[0].runs:
        set_font(run, size=9)

add_section_heading(doc, "DỰ ÁN NỔI BẬT")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("PulseTech  Nền tảng thương mại điện tử thiết bị công nghệ")
set_font(r, size=12, bold=True, color="000000")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Personal Project  |  Full Stack Developer  |  2026")
set_font(r, size=9.5, bold=True, color="1F4E79")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
add_hyperlink(p, "Frontend", "https://github.com/VuIceTea/PulseTech")
r = p.add_run("  |  ")
set_font(r, size=9)
add_hyperlink(p, "Admin", "https://github.com/VuIceTea/PulseTech-Admin")
r = p.add_run("  |  ")
set_font(r, size=9)
add_hyperlink(p, "Backend", "https://github.com/VuIceTea/PulseTech-Backend")

add_bullet(doc, "Xây dựng hệ thống full stack gồm customer storefront, admin dashboard và backend microservices với một API Gateway làm điểm truy cập thống nhất.")
add_bullet(doc, "Phát triển bốn module Maven: API Gateway, Product Service, Auth Service và Order Service; tổ chức dữ liệu trên MongoDB Atlas và giao tiếp nội bộ giữa các service.")
add_bullet(doc, "Thiết kế danh mục sản phẩm hỗ trợ màu sắc, dung lượng, mức giá chênh lệch và tồn kho độc lập cho từng biến thể; kiểm tra tồn kho trước khi tạo đơn và giảm tồn kho đúng phiên bản.")
add_bullet(doc, "Xây dựng hành trình mua hàng từ tìm kiếm, lọc sản phẩm, wishlist, giỏ hàng, voucher, checkout đến lịch sử và theo dõi trạng thái đơn hàng.")
add_bullet(doc, "Tích hợp COD, VNPay, MoMo và Stripe; xử lý URL thanh toán, callback, xác minh giao dịch và cập nhật trạng thái đơn hàng.")
add_bullet(doc, "Phát triển voucher theo khách hàng, giới hạn lượt dùng và hai nhóm giảm giá sản phẩm hoặc phí vận chuyển; hỗ trợ áp dụng đồng thời một mã ở mỗi nhóm.")
add_bullet(doc, "Xây dựng dashboard quản trị sản phẩm, tồn kho, đơn hàng, khách hàng và voucher; hỗ trợ xem chi tiết đơn hàng, kiểm soát luồng trạng thái, khóa tài khoản và gán role.")
add_bullet(doc, "Tích hợp Gemini AI để tư vấn dựa trên danh mục, giá, tồn kho từng biến thể và chính sách được tải trực tiếp từ backend thay vì dữ liệu tĩnh.")
add_bullet(doc, "Container hóa backend bằng Docker multi-stage build và Docker Compose; triển khai frontend qua Vercel và cấu hình backend theo mô hình dịch vụ độc lập.")

add_section_heading(doc, "KIẾN TRÚC VÀ PHẠM VI")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run(
    "Hai ứng dụng Next.js gọi API qua đường dẫn proxy /backend-api. Spring Cloud Gateway định tuyến request "
    "đến Product Service, Auth Service hoặc Order Service. Hệ thống hiện có khoảng 140 file mã nguồn "
    "TypeScript và Java, tương đương gần 14.000 dòng mã ứng dụng, chưa tính CSS, cấu hình và tài nguyên."
)
set_font(r, size=9.7)

add_section_heading(doc, "ĐIỂM KỸ THUẬT TIÊU BIỂU")
add_bullet(doc, "Quản lý state phía client bằng React Context cho authentication, cart và wishlist; đồng bộ giỏ hàng khách với tài khoản sau khi đăng nhập.")
add_bullet(doc, "Xây dựng API cho sản phẩm, nội dung website, tài khoản, địa chỉ, wishlist, giỏ hàng, voucher, đơn hàng và thanh toán.")
add_bullet(doc, "Áp dụng validation cho dữ liệu đầu vào và ràng buộc chuyển trạng thái đơn hàng theo đúng tiến trình nghiệp vụ.")
add_bullet(doc, "Sử dụng skeleton loading, toast notification, responsive layout và giao diện quản trị có chế độ sáng tối.")
add_bullet(doc, "Viết kiểm thử đơn vị cho truy vấn, cập nhật sản phẩm, khởi tạo dữ liệu và DTO đơn hàng bằng Spring Boot Test.")

add_section_heading(doc, "NGÔN NGỮ")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Tiếng Việt: Bản ngữ    |    Tiếng Anh: Đọc hiểu tài liệu kỹ thuật")
set_font(r, size=9.7)

for section in doc.sections:
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(4)
    fr = fp.add_run("Nguyễn Phi Vũ  |  Full Stack Developer  |  PulseTech")
    set_font(fr, size=8, color="777777")

props = doc.core_properties
props.title = "Nguyễn Phi Vũ Full Stack Developer"
props.subject = "CV dự án PulseTech"
props.author = "Nguyễn Phi Vũ"
props.keywords = "Full Stack Developer, Next.js, React, Java, Spring Boot, Microservices, MongoDB, Docker"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
